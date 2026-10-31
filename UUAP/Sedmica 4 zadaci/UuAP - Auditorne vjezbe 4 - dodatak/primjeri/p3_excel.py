# =============================================================
# PRIMJER 3: Excel fajlovi u pandasu (import, export, Excel → CSV)
# Uvod u analizu podataka - Vježbe, Sedmica 4
# =============================================================
# Instalacija: pip install pandas openpyxl
#   (openpyxl je biblioteka koju pandas koristi za .xlsx fajlove)
#
# studenti.xlsx ima DVA sheeta: "Grupa_A" (20 studenata)
# i "Grupa_B" (20 studenata).
# =============================================================
import csv
import os
import pandas as pd

pd.set_option("display.max_columns", None)
pd.set_option("display.width", 120)

FOLDER = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.path.join(FOLDER, "studenti.xlsx")

# --- 1. Učitavanje jednog sheeta -----------------------------
# Bez sheet_name pandas učitava samo PRVI sheet.
grupa_a = pd.read_excel(EXCEL, sheet_name="Grupa_A")
print("Grupa A:", grupa_a.shape)
print(grupa_a.head(3))

# --- 2. Svi sheetovi odjednom: sheet_name=None ---------------
# Dobijemo RJEČNIK: ključ = naziv sheeta, vrijednost = DataFrame
svi_sheetovi = pd.read_excel(EXCEL, sheet_name=None)
print("\nSheetovi u fajlu:", list(svi_sheetovi.keys()))

# --- 3. Spajanje sheetova: pd.concat -------------------------
svi = pd.concat([svi_sheetovi["Grupa_A"], svi_sheetovi["Grupa_B"]], ignore_index=True)
print("Nakon concat:", svi.shape)          # (40, 7)
print(svi.tail(3))

# --- 4. Izmjena i export u Excel: to_excel -------------------
svi["prosjek"] = ((svi["matematika"] + svi["programiranje"] + svi["statistika"]) / 3).round(2)
IZLAZ_XLSX = os.path.join(FOLDER, "studenti_svi.xlsx")
svi.to_excel(IZLAZ_XLSX, index=False)     # index=False: ne piši kolonu 0, 1, 2...
print("\nSačuvano:", IZLAZ_XLSX)          # postojeći fajl se PREPIŠE

# --- 5. Excel → CSV, metoda 1: pandas ------------------------
IZLAZ_CSV = os.path.join(FOLDER, "studenti_svi.csv")
svi.to_csv(IZLAZ_CSV, index=False)        # header=True je podrazumijevano
nazad = pd.read_csv(IZLAZ_CSV)
print("CSV učitan nazad:", nazad.shape)

# --- 6. Excel → CSV, metoda 2: openpyxl + csv ----------------
# Bez pandasa: čitamo ćeliju po ćeliju.
# (Postoji i metoda sa bibliotekom xlrd, ali ona radi samo za
#  stari format .xls - za .xlsx koristimo openpyxl.)
from openpyxl import load_workbook

knjiga = load_workbook(EXCEL)
sheet = knjiga["Grupa_A"]                 # ili knjiga.active za aktivni sheet
IZLAZ_A = os.path.join(FOLDER, "grupa_a_openpyxl.csv")
with open(IZLAZ_A, "w", newline="", encoding="utf-8") as f:
    pisac = csv.writer(f)
    for red in sheet.iter_rows(values_only=True):   # red = tuple vrijednosti
        pisac.writerow(red)
print("openpyxl → CSV:", IZLAZ_A)
print("Prva 2 reda tog CSV-a:")
with open(IZLAZ_A, encoding="utf-8") as f:
    print(f.readline().strip())
    print(f.readline().strip())
